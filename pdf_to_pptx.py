#!/usr/bin/env python3
"""
Convert Beamer main.pdf to PPTX with high-res slide images and injected presenter notes.
Outputs:
  - out/[서강대]_알파퀀트_발표자료.pptx
  - out/[서강대]_알파퀀트_발표자료.pdf
"""

import json
import os
import shutil
import tempfile
from pathlib import Path
import pymupdf
from pptx import Presentation
from pptx.util import Inches

def main():
    root = Path(__file__).resolve().parent
    pdf_path = root / "slides" / "main.pdf"
    notes_path = root / "out" / "script_notes.json"
    out_dir = root / "out"
    out_dir.mkdir(parents=True, exist_ok=True)

    pptx_path = out_dir / "[서강대]_알파퀀트_발표자료.pptx"
    pdf_out_path = out_dir / "[서강대]_알파퀀트_발표자료.pdf"

    print(f"Loading PDF from {pdf_path}...")
    doc = pymupdf.open(pdf_path)
    total_pages = len(doc)
    print(f"Total pages: {total_pages}")

    with open(notes_path, "r", encoding="utf-8") as f:
        notes_dict = json.load(f)

    # Initialize PPTX with 16:9 widescreen dimensions
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)
        for i in range(total_pages):
            page_num = i + 1
            page = doc[i]

            # Render at 250 dpi for crisp text & math rendering
            pix = page.get_pixmap(dpi=250)
            img_file = tmp_path / f"slide_{page_num:02d}.png"
            pix.save(str(img_file))

            # Create slide and add image
            slide = prs.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                str(img_file),
                Inches(0),
                Inches(0),
                width=prs.slide_width,
                height=prs.slide_height
            )

            # Inject presenter note
            note_key = str(page_num)
            if note_key in notes_dict:
                note_text = notes_dict[note_key]
                notes_slide = slide.notes_slide
                text_frame = notes_slide.notes_text_frame
                text_frame.text = note_text

            print(f"Rendered slide {page_num}/{total_pages} with notes ({len(note_text) if note_key in notes_dict else 0} chars)")

        print(f"Saving PPTX to {pptx_path}...")
        prs.save(str(pptx_path))
        print("PPTX saved successfully.")

    print(f"Copying PDF to {pdf_out_path}...")
    shutil.copy2(pdf_path, pdf_out_path)
    print("PDF copy complete.")

    # Check file sizes
    print("\n--- Output Verification ---")
    print(f"PPTX Size: {pptx_path.stat().st_size / (1024*1024):.2f} MB")
    print(f"PDF Size: {pdf_out_path.stat().st_size / (1024*1024):.2f} MB")
    print("Done!")

if __name__ == "__main__":
    main()
